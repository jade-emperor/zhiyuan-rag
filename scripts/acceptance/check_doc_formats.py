"""多格式入库验收:语料装载、格式适配器、格式分派与命名。

**权威契约(改代码前先读这里)**

端口(``ports/document_source.py``)::

    DocumentSource.suffixes -> frozenset[str]          # 这个适配器认识的扩展名
    DocumentSource.read(path) -> str                   # 一份文件 → 带标题层级的 markdown

装载(``adapters/document_loader.py``)::

    load_documents(path) -> LoadedCorpus               # 文件或目录;目录递归
    LoadedCorpus.documents -> tuple[Document, ...]     # 读成功的
    LoadedCorpus.skipped   -> tuple[str, ...]          # 不认识的扩展名(相对路径)
    LoadedCorpus.empty     -> tuple[str, ...]          # 认了但没抽出文字(扫描件/空文档)

四条口径:

1. **命名只有一处**:``doc_id`` = 相对路径、``source`` = 顶层目录名去掉 ``_data``、
   单份文件时 ``source`` 取它所在目录。这三种规则对每种格式一样,所以不在适配器里各写一遍;
2. **顺序确定**:按相对路径排序 —— ``chunk_id`` 的幂等建立在「同一份输入两次同样的顺序」上;
3. **不认识的文件只报告不抛错**(``skipped``),但**解析出错必须抛**:少入库一份文档比报错难查;
4. **空文本进 ``empty`` 而不是 documents**:空文档切不出块,混进去只会让块数对不上。

覆盖的格式与它们的标题来源:``.md``/``.txt`` 原文即结构、``.docx`` Word 的 Heading 样式、
``.pptx`` 每页标题、``.pdf`` 书签目录。**不覆盖**:HTML / Excel / 扫描件 OCR(见 ``deferred.md`` §1.4)。

**为什么不用 unstructured**(实测结论,别再往回改):``partition_docx`` 会调 spacy 分句,
而它不装模型、用的时候现从 GitHub 拉 ``en_core_web_sm`` —— 离线环境直接卡在 socket 上;
且它给的是靠大写比例猜出来的标题,而 Word 文档里本来就有 Heading 样式。
详见 ``adapters/office_files.py`` 的模块 docstring。

**端口符合性由 pyright 证明**:``_CONFORMANCE`` 那一行的注解就是断言 —— 适配器少一个方法,
``uv run pyright`` 会红。运行期的 isinstance 对 Protocol 不做结构检查,写了也是假的。
"""

from __future__ import annotations

import ast
import subprocess
import sys
import tempfile
from contextlib import contextmanager
from pathlib import Path

from project_rag.adapters import (
    LoadedCorpus,
    OfficeFileSource,
    PdfFileSource,
    TextFilesSource,
    load_documents,
    supported_suffixes,
)
from project_rag.adapters._text_blocks import join_wrapped
from project_rag.adapters.office_files import _heading_level
from project_rag.ports import DocumentSource

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "src" / "project_rag"

# 端口符合性断言(见模块 docstring):pyright 会读这一行的注解
_CONFORMANCE: tuple[DocumentSource, ...] = (
    TextFilesSource(),
    OfficeFileSource(),
    PdfFileSource(),
)

EXPECTED_SUFFIXES = {".md", ".txt", ".docx", ".pptx", ".pdf"}

# 真实语料(本地才有;缺了就 SKIP 并点名 owner)
DOC_CORPUS = ROOT / "data" / "corpus" / "doc_data"
PDF_CORPUS = ROOT / "data" / "corpus" / "pdf_data"

MARKDOWN_HEAD = """# 潜艇

> 来源:某飞书文档

## 一、Python

### 1. 编译型与解释型语言区别

编译型语言采用提前编译。
"""

from _harness import check, report, SkipCheck


# ── 夹具:现造四种格式,不依赖本地语料 ─────────────────────────────────────────
def _write_docx(path: Path) -> None:
    """一份 docx:Title + Heading 1 + Heading 2 + 正文 + 表格(4×3 那张的缩小版)。"""
    from docx import Document

    document = Document()
    document.add_heading("员工手册", level=0)  # level=0 = Title 样式
    document.add_heading("一、总则", level=1)
    document.add_paragraph("本手册适用于全体员工。")
    document.add_heading("1.1 适用范围", level=2)
    document.add_paragraph("包含试用期员工与实习生。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "客户等级"
    table.cell(0, 1).text = "账期"
    table.cell(1, 0).text = "A 战略客户"
    table.cell(1, 1).text = "月结 60 天"
    document.save(str(path))


def _write_pptx(path: Path) -> None:
    """一份 pptx:第一页有标题,第二页用 Blank 版式(没有标题占位符)。"""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[5])  # Title Only
    title = first.shapes.title  # 占位符可能缺失,类型上是 Shape | None
    if title is None:
        raise AssertionError("Title Only 版式应当有标题占位符 —— 夹具没造对")
    title.text = "第一页标题"
    first.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = (
        "第一页正文"
    )
    second = presentation.slides.add_slide(presentation.slide_layouts[6])  # Blank
    second.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = (
        "第二页正文"
    )
    presentation.save(str(path))


def _write_pdf(path: Path, text: str | None, toc: list[list] | None = None) -> None:
    """一份两页 PDF:``text=None`` 造一份没有文字层的(模拟扫描件)。"""
    import pymupdf

    document = pymupdf.open()
    for number in range(1, 3):
        page = document.new_page()
        if text:
            page.insert_text((72, 100), f"{text} page {number}")
    if toc:
        document.set_toc(toc)
    document.save(str(path))
    document.close()


@contextmanager
def fixtures():
    """造一个混合语料目录 ``handbook_data/``:四种能读的格式 + 一个不认识的 + 一个空的。"""
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp) / "handbook_data"
        root.mkdir()
        (root / "note.md").write_text(MARKDOWN_HEAD, encoding="utf-8-sig")
        _write_docx(root / "handbook.docx")
        _write_pptx(root / "slides.pptx")
        _write_pdf(
            root / "paper.pdf",
            "chapter text",
            toc=[[1, "Chapter One", 1], [2, "Section A", 2]],
        )
        _write_pdf(root / "blank.pdf", None)  # 没有文字层 → empty
        (root / "archive.zip").write_bytes(b"PK\x03\x04")
        yield root


# ── 1. 端口与注册表 ─────────────────────────────────────────────────────────
@check("端口：DocumentSource 声明 suffixes + read,三个适配器结构上符合它,且从包入口导出")
def _() -> str:
    import project_rag.ports as ports_module

    if not hasattr(ports_module, "DocumentSource"):
        raise AssertionError(
            "project_rag.ports 里没有 DocumentSource —— ports/__init__.py 的 import 列表"
            "与 __all__ 各要加一行(漏 `__all__` 会让 `from ... import *` 少一个名字)"
        )
    tree = ast.parse((SRC / "ports" / "document_source.py").read_text(encoding="utf-8"))
    klass = next(
        (node for node in tree.body if isinstance(node, ast.ClassDef) and node.name == "DocumentSource"),
        None,
    )
    if klass is None:
        raise AssertionError("ports/document_source.py 里没有 DocumentSource")
    annotated = {
        node.target.id
        for node in klass.body
        if isinstance(node, ast.AnnAssign) and isinstance(node.target, ast.Name)
    }
    methods = {node.name for node in klass.body if isinstance(node, ast.FunctionDef)}
    if "suffixes" not in annotated or "read" not in methods:
        raise AssertionError(
            f"端口声明不全:suffixes={sorted(annotated)} / 方法={sorted(methods)}；"
            "分派靠 suffixes、解析靠 read,缺一个都得在调用方各写一遍"
        )

    for source in _CONFORMANCE:
        name = type(source).__name__
        if not isinstance(type(source).suffixes, frozenset):
            raise AssertionError(f"{name} 的 suffixes 必须是 frozenset(注册表靠它建索引)")
        if not callable(getattr(source, "read", None)):
            raise AssertionError(f"{name} 没有可调用的 read")
    # 符合性靠 `_CONFORMANCE` 那一行的注解让 pyright 断,不要求显式继承端口 ——
    # Protocol 的意义就是结构化子类型;以前那条"必须显式继承"的检查只抓到写法、没抓到 bug
    # （删)。
    return "read + suffixes 声明齐全;三个适配器由 pyright 断符合性"


@check("注册表：SOURCES 覆盖 md/txt/docx/pptx/pdf,且后缀两两不重叠")
def _() -> str:
    expected = EXPECTED_SUFFIXES
    got = set(supported_suffixes())
    if got != expected:
        raise AssertionError(
            f"支持的后缀是 {sorted(got)},契约要求 {sorted(expected)}；"
            "加格式 = 加一个适配器 + 在这一条里改契约(先改契约)"
        )
    seen: dict[str, str] = {}
    for source in _CONFORMANCE:
        for suffix in source.suffixes:
            if suffix in seen:
                raise AssertionError(
                    f"{suffix} 被 {seen[suffix]} 与 {type(source).__name__} 同时认领 —— "
                    "分派会按注册顺序吞掉一个,另一种永远不会被用到"
                )
            seen[suffix] = type(source).__name__
    if len(seen) != len(expected):
        raise AssertionError(f"注册表里只有 {sorted(seen)},少了 {sorted(expected - set(seen))}")
    return "、".join(f"{suffix}←{name}" for suffix, name in sorted(seen.items()))


# ── 2. 文本整理口径(所有格式共用) ────────────────────────────────────────────
@check("文本整理：中文折行不补空格 / 拉丁文折行补空格 / 控制字符与私用区字符被去掉")
def _() -> str:
    if join_wrapped("语言\n模型") != "语言模型":
        raise AssertionError(
            f"中文折行被补了空格({join_wrapped('语言\n模型')!r})—— BM25 那一路会把词切碎"
        )
    if join_wrapped("the\nmodel") != "the model":
        raise AssertionError(
            f"拉丁文折行没补空格({join_wrapped('the\nmodel')!r})—— 会粘成 themodel,搜不回来"
        )
    junk = join_wrapped("\x00\x01\uf8ee")
    if junk != "":
        raise AssertionError(f"字体垃圾没被清掉:{junk!r}(PDF 的嵌入字体常吐这些)")
    mixed = join_wrapped("第一行\n第二行\nthird\nline")
    if mixed != "第一行第二行 third line":
        raise AssertionError(
            f"中英混排的接行判据不对:{mixed!r}(中文接中文不留空格,中文接英文要留)"
        )
    return "中文不补空格、拉丁文补空格、垃圾字符清空"


# ── 3. 各格式适配器 ─────────────────────────────────────────────────────────
@check("纯文本：读原文(标题行原样保留)且吃掉 BOM")
def _() -> str:
    with fixtures() as root:
        text = TextFilesSource().read(root / "note.md")
    if text != MARKDOWN_HEAD:
        raise AssertionError(
            "读出来的正文与原文不一致 —— markdown 的结构本来就在正文里,不该被任何解析器改写"
        )
    if text.startswith("\ufeff"):
        raise AssertionError("BOM 没被吃掉 —— 该用 utf-8-sig(否则第一个标题会变成 \\ufeff#)")
    return f"{len(text)} 字,逐字等于原文"


@check("docx：Title→#、Heading 1→##、Heading 2→###,且表格没被丢掉")
def _() -> str:
    with fixtures() as root:
        text = OfficeFileSource().read(root / "handbook.docx")
    for wanted in ("# 员工手册", "## 一、总则", "### 1.1 适用范围"):
        if wanted not in text:
            raise AssertionError(
                f"没找到 {wanted!r} —— Word 的 Heading 样式没被翻成 markdown 标题。实际:\n{text}"
            )
    if "本手册适用于全体员工。" not in text:
        raise AssertionError("正文段落丢了")
    if "| 客户等级 | 账期 |" not in text or "| --- | --- |" not in text:
        raise AssertionError(
            f"表格没被翻成 markdown 表头 —— Document.paragraphs 不含表格,必须显式收。实际:\n{text}"
        )
    if "A 战略客户" not in text:
        raise AssertionError("表格数据行丢了")
    return "三级标题 + 正文 + 表格都在"


@check("docx 样式名：中文 Word 的「标题 N」与 style_id 两路都认,级别封顶到 6")
def _() -> str:
    class _Style:
        def __init__(self, name: str, style_id: str) -> None:
            self.name, self.style_id = name, style_id

    class _Paragraph:
        def __init__(self, name: str, style_id: str) -> None:
            self.style = _Style(name, style_id)

    class _NoStyle:
        @property
        def style(self):
            raise KeyError("styles.xml 里没有这个样式")

    cases = [
        (_Style("标题 1", "Heading1"), 2),
        (_Style("Heading 3", "Heading3"), 4),
        (_Style("标题 9", "Heading9"), 6),  # 超深标题封顶,不许吐 #######
        (_Style("Title", "Title"), 1),
        (_Style("Normal", "Normal"), 0),
        (_NoStyle(), 0),  # 样式表里没有 → 当正文,不许整篇读不出来
    ]
    for index, (style, want) in enumerate(cases):
        paragraph = _Paragraph("", "")
        paragraph.style = style  # type: ignore[assignment]
        got = _heading_level(paragraph)
        if got != want:
            raise AssertionError(
                f"第 {index + 1} 例:样式 {style!r} 应为级别 {want},实际 {got}。"
                "中文版 Word 的样式名是「标题 1」而 style_id 仍是 Heading1,两路都要认"
            )
    return "标题 1/Heading 3/封顶/Title/正文/缺样式 六例都对"


@check("pptx：每页一个 ## 标题,没有标题的页用「第 N 页」占位,正文在内")
def _() -> str:
    with fixtures() as root:
        text = OfficeFileSource().read(root / "slides.pptx")
    if "## 第一页标题" not in text or "## 第 2 页" not in text:
        raise AssertionError(
            f"每页标题没标出来(占位规则:没有标题就用「第 N 页」)。实际:\n{text}"
        )
    if "第一页正文" not in text or "第二页正文" not in text:
        raise AssertionError(f"正文框的文字丢了。实际:\n{text}")
    return "2 页:标题页 + 占位页,正文都在"


@check("pdf：书签目录翻成 #/## 标题,正文按段落分块(块间空行)")
def _() -> str:
    with fixtures() as root:
        text = PdfFileSource().read(root / "paper.pdf")
    if not text.startswith("# Chapter One"):
        raise AssertionError(
            f"第一页开头的书签(1 级)没被翻成 #。实际开头:{text[:120]!r}"
        )
    if "## Section A" not in text:
        raise AssertionError("第二页的 2 级书签没被翻成 ##")
    if "chapter text page 1" not in text or "chapter text page 2" not in text:
        raise AssertionError(f"正文没抽出来。实际:\n{text}")
    if "\n\n" not in text:
        raise AssertionError("块之间没有空行 —— 切分器找不到段落边界,父块会在任意位置断开")
    return "书签 1/2 级 → #/##,两页正文都在"


@check("pdf 扫描件：没有文字层时给空串(不抛错),由装载层报进 empty")
def _() -> str:
    with fixtures() as root:
        text = PdfFileSource().read(root / "blank.pdf")
    if text.strip():
        raise AssertionError(f"没有文字层的 PDF 应给空串,实际 {text[:80]!r}")
    return "空串(不是异常)"


# ── 4. 装载:分派、命名、排序、两份清单 ──────────────────────────────────────
@check("装载：混合目录按后缀分派,source 去掉 _data,按相对路径排序")
def _() -> str:
    with fixtures() as root:
        loaded = load_documents(root)
    if not isinstance(loaded, LoadedCorpus):
        raise AssertionError("load_documents 必须返回 LoadedCorpus(三份清单各有各的修法)")
    ids = [document.doc_id for document in loaded.documents]
    want = ["handbook.docx", "note.md", "paper.pdf", "slides.pptx"]
    if ids != want:
        raise AssertionError(f"文档顺序/分派不对:得到 {ids},期望 {want}(按相对路径排序)")
    for document in loaded.documents:
        if document.source != "handbook":
            raise AssertionError(
                f"{document.doc_id} 的 source 应为 'handbook'(handbook_data 去掉 _data),"
                f"实际 {document.source!r}"
            )
    if loaded.skipped != ("archive.zip",):
        raise AssertionError(f"不认识的 .zip 该进 skipped,实际 {loaded.skipped}")
    if loaded.empty != ("blank.pdf",):
        raise AssertionError(f"没有文字的 PDF 该进 empty,实际 {loaded.empty}")
    return f"{len(ids)} 份文档:{ids};skipped={loaded.skipped};empty={loaded.empty}"


@check("装载：单份文件、不存在的路径、确定性重跑")
def _() -> str:
    with fixtures() as root:
        single = load_documents(root / "note.md")
        if [document.doc_id for document in single.documents] != ["note.md"]:
            raise AssertionError(f"单份文件的 doc_id 应是文件名,实际 {single.documents}")
        if single.documents[0].source != "handbook":
            raise AssertionError("单份文件的 source 要取它所在目录(去掉 _data),实际不一致")

        first = load_documents(root)
        second = load_documents(root)
        if [d.doc_id for d in first.documents] != [d.doc_id for d in second.documents]:
            raise AssertionError("两次装载顺序不一致 —— chunk_id 的幂等就建立在这个顺序上")
        if [d.text for d in first.documents] != [d.text for d in second.documents]:
            raise AssertionError("两次装载正文不一致 —— 解析里有随机/时间因素")

        try:
            load_documents(root / "不存在")
        except FileNotFoundError:
            pass
        else:
            raise AssertionError("路径不存在必须抛 FileNotFoundError,不能给一个空结果")
    return "单文件命名 / 两次一致 / 不存在即报错"


@check("CLI：ingest --path <混合目录> --dry-run 跑通,并报出不认识的文件")
def _() -> str:
    with fixtures() as root:
        proc = subprocess.run(
            [sys.executable, "-m", "project_rag", "ingest", "--path", str(root), "--dry-run"],
            capture_output=True,
            text=True,
            timeout=300,
            cwd=str(ROOT),
        )
    if proc.returncode != 0:
        raise AssertionError(
            f"退出码 {proc.returncode}；stderr={proc.stderr[-300:]!r}。"
            "dry-run 不该需要 Milvus / 模型 —— 它只加载与切分"
        )
    if "切分:" not in proc.stdout or "加载 4 份文档" not in proc.stdout:
        raise AssertionError(f"dry-run 没打印加载与切分统计,实际 stdout={proc.stdout[-300:]!r}")
    if "archive.zip" not in proc.stderr:
        raise AssertionError(
            f"不认识的文件没被告知(stderr={proc.stderr[-200:]!r})—— "
            "少入库一份文档必须让人当场看见"
        )
    summary = [line for line in proc.stdout.splitlines() if line.startswith("切分:")][0]
    return summary


# ── 5. 真实语料(本地才有) ──────────────────────────────────────────────────
@check("真实语料：doc_data 的每份 docx 都读出标题,pdf_data 的每份 PDF 都抽出正文")
def _() -> str:
    if not DOC_CORPUS.exists() or not PDF_CORPUS.exists():
        raise SkipCheck(
            f"真实语料目录不在本地（{DOC_CORPUS.name} / {PDF_CORPUS.name}）；"
            "owner：数据装载（见 data/README.md）"
        )
    office = load_documents(DOC_CORPUS)
    pdfs = load_documents(PDF_CORPUS)
    if not office.documents or not pdfs.documents:
        raise AssertionError(
            f"真实语料没读出文档：docx {len(office.documents)} 份 / PDF {len(pdfs.documents)} 份"
        )
    for document in office.documents:
        if "#" not in document.text:
            raise AssertionError(
                f"{document.doc_id} 里一个 markdown 标题都没有 —— Word 的 Heading 样式没被翻出来"
            )
    for document in pdfs.documents:
        if len(document.text) < 10_000:
            raise AssertionError(
                f"{document.doc_id} 只抽出 {len(document.text)} 字 —— 文字层没读到"
                "(扫描件会是这样,这两份不是)"
            )
        if "\n#" not in document.text:
            raise AssertionError(f"{document.doc_id} 没有书签标题 —— 目录没被用上")
    sizes = "、".join(
        f"{document.doc_id}={len(document.text)}字" for document in pdfs.documents
    )
    return f"docx {len(office.documents)} 份有标题；PDF {len(pdfs.documents)} 份:{sizes}"


def main() -> int:
    return report("多格式入库验收（装载 / 分派 / 命名）")


if __name__ == "__main__":
    sys.exit(main())
