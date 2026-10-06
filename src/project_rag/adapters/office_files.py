"""
Office 文档来源:``.docx`` / ``.pptx``(python-docx / python-pptx)。

**为什么不用 ``unstructured``**:它装是装了,但 ``partition_docx`` 在切分元素时会调
``is_possible_narrative_text`` → ``exceeds_cap_ratio`` → ``sentence_count`` → spacy 分句,
而它**不装模型、用的时候现下**:``unstructured/nlp/tokenize.py`` 里写死从 GitHub 拉
``en_core_web_sm``(120s 超时 + 哈希校验)。离线环境里这条路径直接挂住 —— 实测「读一个 37KB 的
docx」卡在 socket 上五分钟没动静。而且它给的是**猜出来的**标题(靠大写比例/句子数启发式),
Word 文档里明明有现成的 Heading 样式。

直接读样式的结果严格更好:

* 标题层级是**文档自己声明的**(``Heading 1`` / ``标题 1`` / ``Title`` 样式),不是猜的;
* 表格要显式收集 —— ``Document.paragraphs`` **不含表格里的文字**,只拼段落会静默丢表
  (``customer_credit_v1.docx`` 那张 4×3 的授信表就是这么丢的);
* 零新增依赖、离线可用、快(4 份 docx 共 0.1s)。

代价说清楚:文本框 / 页眉页脚 / 批注里的文字读不到(它们在 body 之外)。
这些位置对我们的语料没有信息量,真要收再说。
"""

from __future__ import annotations

import re
from collections.abc import Iterable, Iterator, Sequence
from pathlib import Path
from typing import ClassVar

from ._text_blocks import MAX_HEADING_LEVEL, join_wrapped

# Word 的标题样式:英文 "Heading 1"、中文 "标题 1";style_id 一般是 "Heading1"(与界面语言无关)
_HEADING_RE = re.compile(r"^(?:heading|标题)\s*(\d+)$", re.IGNORECASE)


def _table_to_markdown(rows: Iterable[Sequence[str]]) -> str:
    """输入:表格的行列文本。行为:拼成 markdown 表格。输出:markdown 文本(空表给空串)。

    单元格里的换行折成空格、``|`` 转义 —— 不折的话表格会被自己的换行拆成两半,
    竖线不转义则会让列数对不上(markdown 表格靠竖线认列)。
    """
    lines: list[list[str]] = [
        [join_wrapped(cell).replace("|", "\\|") for cell in row] for row in rows
    ]
    lines = [row for row in lines if any(cell for cell in row)]
    if not lines:
        return ""
    width = max(len(row) for row in lines)
    padded = [row + [""] * (width - len(row)) for row in lines]
    head, *body = padded
    out = ["| " + " | ".join(head) + " |", "| " + " | ".join(["---"] * width) + " |"]
    out += ["| " + " | ".join(row) + " |" for row in body]
    return "\n".join(out)


def _heading_level(paragraph) -> int:
    """输入:一个 docx 段落。行为:看它的样式名是不是标题。输出:标题级别(0 = 不是标题)。

    样式名与 style_id 都要看:中文版 Word 的样式名叫「标题 1」而 style_id 仍是 ``Heading1``,
    自己另存过的文档又可能反过来。取不到样式就当正文 —— 宁可少一个标题,不可整篇读不出来。
    """
    try:
        names = (paragraph.style.name or "", paragraph.style.style_id or "")
    except (KeyError, AttributeError):  # 样式表里没有这个样式:按正文处理
        return 0
    for name in names:
        if name.strip().lower() == "title":
            return 1
        matched = _HEADING_RE.match(name.strip())
        if matched:
            # Title 用 #,Heading 1 用 ## —— 让文档标题与章节标题不在同一层
            return min(int(matched.group(1)) + 1, MAX_HEADING_LEVEL)
    return 0


def _docx_blocks(path: Path) -> Iterator[str]:
    """输入:一份 docx。行为:按**正文顺序**遍历段落与表格。输出:块文本序列。

    为什么不用 ``Document.paragraphs``:它只给段落,表格被整个跳过,而且段落与表格的
    前后顺序会丢(段落全排在表格前面)。所以走 body 的元素序列,遇到什么解什么。
    """
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(element, document)
            text = join_wrapped(paragraph.text)
            if not text:
                continue
            level = _heading_level(paragraph)
            yield f"{'#' * level} {text}" if level else text
        elif tag == "tbl":
            table = Table(element, document)
            markdown = _table_to_markdown(
                [cell.text for cell in row.cells] for row in table.rows
            )
            if markdown:
                yield markdown


def _pptx_blocks(path: Path) -> Iterator[str]:
    """输入:一份 pptx。行为:一页一页地取标题、正文与表格。输出:块文本序列。

    每页标题定成 ``##``:一页就是一段独立的上下文,标题丢了整页文字会跟前后页粘成一块。
    没有标题的页用「第 N 页」占位 —— 位置信息对「在第几页讲过」这类问题有用。
    """
    from pptx import Presentation
    from pptx.shapes.autoshape import Shape as AutoShape
    from pptx.shapes.graphfrm import GraphicFrame

    presentation = Presentation(str(path))
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = ""
        if title_shape is not None and title_shape.has_text_frame:
            title = join_wrapped(title_shape.text_frame.text)
        yield f"## {title or f'第 {number} 页'}"

        for shape in slide.shapes:
            if shape is title_shape:
                continue
            # 表格是 GraphicFrame,文本框是 Shape —— 都得判类,不能靠 ``has_table`` 那种
            # 「有没有这个属性」的鸭子判断:``BaseShape`` 上根本没有 ``text_frame``。
            if isinstance(shape, GraphicFrame) and shape.has_table:
                markdown = _table_to_markdown(
                    [[cell.text for cell in row.cells] for row in shape.table.rows]
                )
                if markdown:
                    yield markdown
            elif isinstance(shape, AutoShape):
                text = join_wrapped(shape.text_frame.text)
                if text:
                    yield text


class OfficeFileSource:
    """``.docx`` / ``.pptx`` → 带标题层级的 markdown。"""

    suffixes: ClassVar[frozenset[str]] = frozenset({".docx", ".pptx"})

    def read(self, path: Path) -> str:
        """输入:一份 .docx/.pptx 的路径。行为:按后缀选解析器并抽出正文。输出:markdown 文本。

        块之间空一行:切分器按 ``\\n\\n`` 找段落边界,挤在一起会让标题和正文粘成一块
        (父块就切不开了)。
        """
        blocks = _docx_blocks if path.suffix.lower() == ".docx" else _pptx_blocks
        return "\n\n".join(blocks(path))
